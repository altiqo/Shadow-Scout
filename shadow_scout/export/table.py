"""Табличные экспортёры: XLSX, CSV, Markdown, JSON."""

from __future__ import annotations

import csv
import json
from pathlib import Path

from shadow_scout.export.common import COLUMNS, rows_for
from shadow_scout.models import VERDICT_COLORS, SearchResult


def write_csv(result: SearchResult, path: Path) -> None:
    rows = rows_for(result)
    with path.open("w", encoding="utf-8-sig", newline="") as fh:
        writer = csv.writer(fh, delimiter=";")
        writer.writerow([title for _, title in COLUMNS])
        for row in rows:
            writer.writerow([row[key] for key, _ in COLUMNS])


def write_json(result: SearchResult, path: Path) -> None:
    payload = result.model_dump(mode="json")
    payload["reports"] = [r.model_dump(mode="json") for r in result.sorted_reports()]
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def write_markdown(result: SearchResult, path: Path) -> None:
    rows = rows_for(result)
    keys = ["rank", "provider", "countries", "asns", "ipv4", "blocked_share", "ip_type", "trial", "hourly", "price", "survivability", "verdict", "confidence"]
    titles = {k: t for k, t in COLUMNS}
    lines = [f"# {result.title}", "", f"Создан: {result.generated_at.strftime('%Y-%m-%d %H:%M UTC')}  ", f"Запрос: {result.query.describe()}", ""]
    lines.append("| " + " | ".join(titles[k] for k in keys) + " |")
    lines.append("|" + "|".join("---" for _ in keys) + "|")
    for row in rows:
        lines.append("| " + " | ".join(row[k].replace("|", "/") for k in keys) + " |")
    lines.append("")
    for report in result.sorted_reports():
        p = report.provider
        lines.append(f"## {p.name} — выживаемость {report.survivability:.0f}/100 ({report.verdict})")
        if p.website:
            lines.append(f"Сайт: {p.website}  ")
        for a in report.asns:
            lines.append(f"- AS{a.asn} — {a.holder or '?'}; IPv4: {a.ipv4_count}; в списках: {a.blocked_share:.2%}; префиксов: {len(a.prefixes_v4)}")
        if report.hard_flags:
            lines.append("- **Флаги:** " + "; ".join(report.hard_flags))
        lines.append("")
        lines.append("| Сигнал | Риск | Описание |")
        lines.append("|---|---|---|")
        for s in report.signals:
            risk = f"{s.risk * 100:.0f}" if s.risk is not None else "—"
            lines.append(f"| {s.title} | {risk} | {s.summary.replace('|', '/')} |")
        lines.append("")
        if report.recommendations:
            lines.append("Рекомендации:")
            lines.extend(f"- {r}" for r in report.recommendations)
        if report.warnings:
            lines.append("Предупреждения:")
            lines.extend(f"- {w}" for w in report.warnings)
        lines.append("")
    lines.append("## Источники")
    for s in result.sources:
        lines.append(f"- {'✅' if s.ok else '⚠️'} {s.name}: {s.detail}")
    path.write_text("\n".join(lines), encoding="utf-8")


def write_xlsx(result: SearchResult, path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Провайдеры"
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(bold=True, color="FFFFFF")
    thin = Side(style="thin", color="D1D5DB")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append([title for _, title in COLUMNS])
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border
    reports = result.sorted_reports()
    for rank, (row, report) in enumerate(zip(rows_for(result), reports, strict=True), start=2):
        values = []
        for key, _ in COLUMNS:
            val = row[key]
            if key in ("survivability", "risk"):
                val = float(val)
            elif key == "blocked_share" and val not in ("—", ""):
                val = float(val)
            elif key == "price" and val not in ("—", ""):
                val = float(val)
            values.append(val)
        ws.append(values)
        color = VERDICT_COLORS.get(report.verdict, "#94a3b8").lstrip("#")
        verdict_col = [k for k, _ in COLUMNS].index("verdict") + 1
        ws.cell(row=rank, column=verdict_col).fill = PatternFill("solid", fgColor=color)
        ws.cell(row=rank, column=verdict_col).font = Font(bold=True, color="FFFFFF" if report.verdict in ("critical", "high") else "111827")
        for cell in ws[rank]:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = {"provider": 26, "website": 30, "countries": 34, "asns": 16, "holder": 34, "flags": 40, "verdict": 18}
    for idx, (key, _) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = widths.get(key, 13)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    details = wb.create_sheet("Сигналы")
    details.append(["Провайдер", "Сигнал", "Риск (0-100)", "Вес", "Статус", "Описание"])
    for cell in details[1]:
        cell.fill = header_fill
        cell.font = header_font
    for report in reports:
        for s in report.signals:
            details.append([report.provider.name, s.title, round(s.risk * 100) if s.risk is not None else None, s.weight, s.status.value, s.summary])
    for col, width in zip("ABCDEF", (26, 36, 12, 8, 12, 90), strict=True):
        details.column_dimensions[col].width = width

    recs = wb.create_sheet("Рекомендации")
    recs.append(["Провайдер", "Рекомендация"])
    for cell in recs[1]:
        cell.fill = header_fill
        cell.font = header_font
    for report in reports:
        for rec in report.recommendations:
            recs.append([report.provider.name, rec])
    recs.column_dimensions["A"].width = 26
    recs.column_dimensions["B"].width = 120

    src = wb.create_sheet("Источники")
    src.append(["Источник", "Статус", "Детали"])
    for s in result.sources:
        src.append([s.name, "OK" if s.ok else "проблема", s.detail])
    src.column_dimensions["A"].width = 36
    src.column_dimensions["C"].width = 60
    wb.save(path)
