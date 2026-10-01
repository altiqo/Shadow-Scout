"""Отображение «лучших по локациям»: терминал и все форматы экспорта."""

from __future__ import annotations

import io
import json
from pathlib import Path

import pytest
from rich.console import Console

from shadow_scout import selection
from shadow_scout.analysis.engine import AnalysisEngine
from shadow_scout.export import export_many
from shadow_scout.export.common import pick_labels, pick_rows, picked_reports, yes_no
from shadow_scout.models import Location, Provider, SearchQuery
from shadow_scout.ui import views
from shadow_scout.ui.theme import THEME


def make_providers() -> list[Provider]:
    def p(pid, name, asn, cc, **kw):
        base = dict(
            id=pid, name=name, website=f"https://{pid}.example", asns=[asn], hq_country=cc,
            locations=[Location(country=cc, city="Stockholm" if cc == "SE" else "Helsinki")], size="small", popularity_ru=1, vps=True,
        )
        return Provider(**{**base, **kw})

    return [
        p("glesys", "GleSYS", 42708, "SE", hourly=True, trial=False, min_price_eur=5),
        p("tiny", "Tiny ISP", 99999, "SE", hourly=None, source="catalog", vps=None),
        p("upcloud", "UpCloud", 202053, "FI", hourly=True, trial=True),
        p("flokinet", "FlokiNET", 200651, "FI"),
    ]


@pytest.fixture
async def result(services):
    await services.ensure_blocklists()
    providers = make_providers()
    query = SearchQuery(countries=["SE", "FI"], per_location=2)
    plan = selection.plan_candidates(providers, query)
    return await AnalysisEngine(services).run_search(providers, query, "Тест", plan=plan)


def render(renderable, width: int = 200) -> str:
    buffer = io.StringIO()
    Console(file=buffer, width=width, force_terminal=False, color_system=None, theme=THEME).print(renderable)
    return buffer.getvalue()


# ───────────────────────── терминал ─────────────────────────
def test_flag_cell_symbols():
    assert views.flag_cell(True).plain == "✓"
    assert views.flag_cell(True, hint=True).plain == "~"  # определено автоматически
    assert views.flag_cell(False).plain == "—"
    assert views.flag_cell(None).plain == "?"


def test_yes_no_marks_auto_detected_values():
    assert [yes_no(True), yes_no(True, True), yes_no(False), yes_no(None)] == ["да", "да (авто)", "нет", "?"]


async def test_picks_view_groups_by_location(result):
    text = render(views.picks_view(result))
    assert "Швеция (SE)" in text and "Финляндия (FI)" in text
    assert text.index("Швеция (SE)") < text.index("GleSYS") < text.index("Финляндия (FI)") < text.index("UpCloud")
    assert "Stockholm" in text and "AS42708" in text
    assert "[авто]" in text  # каталожная запись помечена
    assert "в базе подошло" in text


def test_picks_view_marks_empty_locations():
    from shadow_scout.models import LocationPick, SearchResult

    result = SearchResult(query=SearchQuery(), picks=[LocationPick(country="HR", candidates=0), LocationPick(country="IS", candidates=3, analyzed=3)])
    text = render(views.picks_view(result))
    assert "Хорватия (HR)" in text and "в базе нет подходящих провайдеров" in text
    assert "Исландия (IS)" in text and "рекомендовать некого" in text


def test_picks_view_without_picks_explains():
    from shadow_scout.models import SearchResult

    assert "Рекомендовать некого" in render(views.picks_view(SearchResult(query=SearchQuery())))


def test_plan_line_mentions_thin_locations():
    providers = [Provider(id="a", name="A", asns=[1], hq_country="IS", locations=[Location(country="IS")], vps=True)]
    plan = selection.plan_candidates(providers, SearchQuery(countries=["IS"]))
    text = views.plan_line(plan).plain
    assert "Кандидатов в базе: 1" in text and "Мало вариантов" in text and "Исландия (IS): 1" in text and "harvest" in text


async def test_results_table_marks_depth_and_unknown_flags(result):
    text = render(views.results_table(result))
    assert "Пров." in text and "●" in text
    assert "?" in text  # у «Tiny ISP» почасовая неизвестна


# ───────────────────────── экспорт ─────────────────────────
async def test_pick_helpers(result):
    labels = pick_labels(result)
    assert labels["glesys"][0].startswith("SE #")
    rows = pick_rows(result)
    assert {r["location"] for r in rows} == {"Швеция (SE)", "Финляндия (FI)"}
    assert all(r["rank"].isdigit() for r in rows)
    ids = [r.provider.id for r in picked_reports(result)]
    assert len(ids) == len(set(ids))


async def test_all_formats_include_locations_section(result, tmp_path: Path):
    paths = {p.suffix[1:]: p for p in export_many(result, ["html", "pdf", "xlsx", "csv", "md", "json"], tmp_path, "r")}
    html = paths["html"].read_text(encoding="utf-8")
    assert "Лучшие по локациям" in html and "Швеция (SE)" in html and "id=\"p-glesys\"" in html and "href='#p-glesys'" in html
    md = paths["md"].read_text(encoding="utf-8")
    assert "## Лучшие по локациям" in md and "### Финляндия (FI)" in md
    assert "Лучший в локациях" in paths["csv"].read_text(encoding="utf-8-sig")
    data = json.loads(paths["json"].read_text(encoding="utf-8"))
    assert data["picks"] and data["candidates_total"] == 4 and {pick["country"] for pick in data["picks"]} == {"SE", "FI"}
    from openpyxl import load_workbook

    wb = load_workbook(paths["xlsx"])
    assert wb.sheetnames[0] == "По локациям" and "Провайдеры" in wb.sheetnames
    sheet = wb["По локациям"]
    assert sheet.cell(1, 1).value == "Локация" and sheet.max_row >= 3
    assert paths["pdf"].stat().st_size > 2000


async def test_export_without_picks_still_works(services, tmp_path: Path):
    """Проверка одного провайдера (без локаций) не должна ломаться на новых разделах."""
    from shadow_scout.models import SearchResult

    await services.ensure_blocklists()
    report = await AnalysisEngine(services).analyze(make_providers()[0])
    single = SearchResult(query=SearchQuery(), reports=[report])
    for p in export_many(single, ["html", "pdf", "xlsx", "csv", "md", "json"], tmp_path, "single"):
        assert p.exists() and p.stat().st_size > 500
    assert "Лучшие по локациям" not in (tmp_path / "single.html").read_text(encoding="utf-8")


async def test_quick_reports_say_holder_was_not_requested(services):
    """У быстрой проверки держатель ASN не запрашивался — в отчётах это не должно выглядеть как «неизвестно»."""
    from shadow_scout.export.common import holder_label

    await services.ensure_blocklists()
    engine = AnalysisEngine(services)
    quick = await engine.analyze(make_providers()[0], mode="quick")
    full = await engine.analyze(make_providers()[0], mode="full")
    assert quick.asns[0].holder is None and holder_label(quick, None) == "не запрашивался (быстрая проверка)"
    assert holder_label(full, full.asns[0].holder) == "GLESYS-AS" and holder_label(full, None) == "?"
